import os
import base64
import requests
from collections import Counter
from datetime import datetime, timezone
from dotenv import load_dotenv
from tqdm import tqdm
from openpyxl import Workbook, load_workbook

# --------------------------------------------------
# Load environment variables
# --------------------------------------------------
load_dotenv()

CLIENT_ID = os.getenv("EBAY_CLIENT_ID")
CLIENT_SECRET = os.getenv("EBAY_CLIENT_SECRET")

if not CLIENT_ID or not CLIENT_SECRET:
    raise RuntimeError("Missing eBay credentials")

MARKETPLACE_ID = "EBAY_US"  # marketplace endpoint

# --------------------------------------------------
# OAuth
# --------------------------------------------------
def get_access_token():
    credentials = f"{CLIENT_ID}:{CLIENT_SECRET}"
    encoded = base64.b64encode(credentials.encode()).decode()

    r = requests.post(
        "https://api.ebay.com/identity/v1/oauth2/token",
        headers={
            "Authorization": f"Basic {encoded}",
            "Content-Type": "application/x-www-form-urlencoded",
        },
        data={
            "grant_type": "client_credentials",
            "scope": "https://api.ebay.com/oauth/api_scope",
        },
        timeout=30,
    )

    if r.status_code != 200:
        raise RuntimeError(r.text)

    return r.json()["access_token"]

# --------------------------------------------------
# Discover Moroccan sellers
# --------------------------------------------------
def find_moroccan_sellers(token, keywords, limit=50, max_pages=12):
    sellers = []

    headers = {
        "Authorization": f"Bearer {token}",
        "X-EBAY-C-MARKETPLACE-ID": MARKETPLACE_ID,
    }

    for keyword in tqdm(keywords, desc="🔍 Discovering sellers", unit="keyword"):
        offset = 0
        for _ in range(max_pages):
            r = requests.get(
                "https://api.ebay.com/buy/browse/v1/item_summary/search",
                headers=headers,
                params={
                    "q": keyword,
                    "limit": limit,
                    "offset": offset,
                    "filter": "itemLocationCountry:MA",
                },
                timeout=30,
            )

            if r.status_code != 200:
                break

            items = r.json().get("itemSummaries", [])
            if not items:
                break

            for item in items:
                seller = item.get("seller", {}).get("username")
                if seller:
                    sellers.append(seller)

            offset += limit

    return Counter(sellers)

# --------------------------------------------------
# Get products
# --------------------------------------------------
def get_products(token, seller, keywords, run_date, collected_at, limit=25):
    products = []

    headers = {
        "Authorization": f"Bearer {token}",
        "X-EBAY-C-MARKETPLACE-ID": MARKETPLACE_ID,
    }

    for keyword in keywords:
        r = requests.get(
            "https://api.ebay.com/buy/browse/v1/item_summary/search",
            headers=headers,
            params={
                "q": keyword,
                "limit": limit,
                "filter": f"itemLocationCountry:MA,seller:{seller}",
            },
            timeout=30,
        )

        if r.status_code != 200:
            continue

        for item in r.json().get("itemSummaries", []):
            products.append([
                run_date,
                collected_at,
                item.get("itemCreationDate"),
                seller,
                item.get("title"),
                (item.get("categories") or [{}])[0].get("categoryName"),
                item.get("price", {}).get("value"),
                item.get("price", {}).get("currency"),
                "MOROCCO",
            ])

    return products

# --------------------------------------------------
# Excel helpers
# --------------------------------------------------
def get_workbook(filename):
    if os.path.exists(filename):
        wb = load_workbook(filename)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append([
            "run_date",
            "collected_at",
            "listing_date",
            "seller",
            "title",
            "category",
            "price",
            "currency",
            "market",
        ])
    return wb, ws

# --------------------------------------------------
# Main
# --------------------------------------------------
def main():
    run_date = datetime.now(timezone.utc).date().isoformat()
    collected_at = datetime.now(timezone.utc).isoformat()

    print("🔐 Getting token...")
    token = get_access_token()
    print("Token OK ✅\n")

    # VERY BROAD keywords to ensure volume
    keywords = [
        "phone", "iphone", "samsung", "android", "tablet", "ipad",
        "laptop", "computer", "pc", "gaming", "monitor",
        "electronics", "charger", "cable", "adapter", "power supply",
        "keyboard", "mouse", "ssd", "hard drive",
        "camera", "headphone", "speaker", "microphone",
        "watch", "smartwatch",
        "car", "auto", "spare part", "engine", "battery",
        "clothing", "shoes", "bag",
    ]

    seller_counter = find_moroccan_sellers(
        token,
        keywords,
        limit=50,
        max_pages=12,
    )

    top_sellers = [s for s, _ in seller_counter.most_common(50)]

    filename = "morocco_market_data.xlsx"
    wb, ws = get_workbook(filename)

    total_rows = 0

    for seller in tqdm(top_sellers, desc="📦 Collecting products", unit="seller"):
        rows = get_products(
            token,
            seller,
            keywords,
            run_date,
            collected_at,
            limit=25,
        )
        for row in rows:
            ws.append(row)
        total_rows += len(rows)

    wb.save(filename)

    size_mb = os.path.getsize(filename) / (1024 * 1024)

    print(f"\n✅ Added {total_rows} rows")
    print(f"📊 Excel size: {size_mb:.2f} MB")
    print("Done ✅")

# --------------------------------------------------
if __name__ == "__main__":
    main()
